"""Exportar normas a Word/PDF con formato profesional (párrafos reales, justificado)."""

from __future__ import annotations

import re
from pathlib import Path
import tempfile

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _normalizar(texto: str) -> str:
    """Normaliza texto: une palabras cortadas, colapsa saltos simples, conserva párrafos."""
    t = re.sub(r"([a-záéíóúñ])([áéíóúñ])\s+([a-z]{1,3})\b", r"\1\2\3", texto)
    t = re.sub(r"\n{2,}", "\u2029", t)
    t = re.sub(r"\n+", " ", t)
    t = t.replace("\u2029", "\n\n")
    t = re.sub(r"[ \t]{2,}", " ", t)
    return t.strip()


def exportar_docx(titulo: str, texto: str, numero: str | None = None, fecha: str | None = None) -> str:
    texto = _normalizar(texto)
    doc = Document()
    # Márgenes
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
    # Estilo base: Arial 11, interlineado 1.5
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)

    # Título
    h = doc.add_heading(titulo, level=1)
    if numero:
        doc.add_paragraph(f"Norma N° {numero}")
    if fecha:
        doc.add_paragraph(f"Publicada: {fecha}")

    # Cuerpo: dividir en párrafos reales, justificados con sangría de primera línea
    parrafos = [p.strip() for p in texto.split("\n") if p.strip()]
    for p in parrafos:
        # Encabezados de fuente/links al final, en tamaño menor
        if p.startswith(("Fuente oficial:", "Links oficiales", "—")) or "https://" in p and len(p) < 200:
            par = doc.add_paragraph(p)
            par.runs[0].font.size = Pt(9)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            par = doc.add_paragraph(p)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.first_line_indent = Cm(1.0)

    doc.add_paragraph()
    par = doc.add_paragraph("Fuente oficial: https://www.bcn.cl/leychile/navegar?idNorma=<id>")
    par.runs[0].font.size = Pt(9)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    return tmp.name


def exportar_pdf(titulo: str, texto: str, numero: str | None = None, fecha: str | None = None) -> str:
    texto = _normalizar(texto)
    # Fallback reportlab (puro Python, sin libs del sistema) — párrafos reales justificados
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_JUSTIFY
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.units import cm

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        doc = SimpleDocTemplate(tmp.name, pagesize=A4, topMargin=2.2*cm, bottomMargin=2.2*cm,
                                leftMargin=2.5*cm, rightMargin=2.5*cm)
        styles = getSampleStyleSheet()
        titulo_style = ParagraphStyle("T2", parent=styles["Heading1"], fontSize=14, leading=18, spaceAfter=10)
        normal = ParagraphStyle("N2", parent=styles["Normal"], fontName="Times-Roman", fontSize=11,
                                leading=16.5, alignment=TA_JUSTIFY, firstLineIndent=1.0*cm, spaceAfter=8)
        meta = ParagraphStyle("M2", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=6)

        story = [Paragraph(_esc(titulo), titulo_style)]
        meta_line = f"Norma N° {numero or ''}" + (f" — Publicada: {fecha}" if fecha else "")
        story.append(Paragraph(_esc(meta_line), meta))
        story.append(Spacer(1, 0.4*cm))

        for p in [p.strip() for p in texto.split("\n") if p.strip()]:
            if p.startswith(("Fuente oficial:", "Links oficiales", "—")) or ("https://" in p and len(p) < 200):
                small = ParagraphStyle("S2", parent=normal, fontSize=9, leading=12, firstLineIndent=0)
                story.append(Paragraph(_esc(p), small))
            else:
                story.append(Paragraph(_esc(p[:2500]), normal))

        story.append(Spacer(1, 0.4*cm))
        small2 = ParagraphStyle("S3", parent=normal, fontSize=9, leading=12, firstLineIndent=0)
        story.append(Paragraph("Fuente oficial: https://www.bcn.cl/leychile/navegar?idNorma=&lt;id&gt;", small2))
        doc.build(story)
        return tmp.name
    except Exception as e:
        return f"PDF no disponible: {e}"


def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
